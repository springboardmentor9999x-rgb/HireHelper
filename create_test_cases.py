import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
wb.remove(wb.active)

test_cases = {
    'Authentication': [
        ('AUTH-001', 'User registration with valid credentials', 'Email: user@example.com, Password: Secure@123!, Confirm: Secure@123!', 'Account created successfully, User redirected to OTP verification page, Success message displayed', '', ''),
        ('AUTH-002', 'User registration with weak password', 'Email: user@example.com, Password: 123456', 'Error message: Password must contain uppercase, lowercase, number, and special character, Form remains on registration page', '', ''),
        ('AUTH-003', 'User registration with existing email', 'Email: existing@example.com, Password: Secure@123!', 'Error message: Email already registered, Registration fails, User redirected to login', '', ''),
        ('AUTH-004', 'OTP verification with correct OTP', 'Email: user@example.com, OTP: 123456 (sent to email)', 'Account verified successfully, User redirected to login page, Success notification displayed', '', ''),
        ('AUTH-005', 'OTP verification with incorrect OTP', 'Email: user@example.com, OTP: 999999', 'Error message: Invalid OTP, OTP field cleared, Retry option available, Resend OTP link displayed', '', ''),
        ('AUTH-006', 'User login with correct credentials', 'Email: verified@example.com, Password: Secure@123!', 'User logged in successfully, Dashboard loaded, Session token generated, User data displayed', '', ''),
        ('AUTH-007', 'User login with incorrect password', 'Email: verified@example.com, Password: WrongPass@123!', 'Error message: Invalid credentials, Login form remains on screen, Password field cleared for security', '', ''),
        ('AUTH-008', 'User login with unverified account', 'Email: unverified@example.com, Password: Secure@123!', 'Error message: Account not verified. Check your email for OTP, Redirect to OTP verification page', '', ''),
        ('AUTH-009', 'User login with non-existent email', 'Email: notregistered@example.com, Password: Secure@123!', 'Error message: No account found with this email, Login form remains on screen', '', ''),
        ('AUTH-010', 'Resend OTP functionality', 'Email: user@example.com, Button: Resend OTP', 'New OTP sent to email, Success message: OTP sent to your email, Timer starts for next resend (60 seconds)', '', ''),
    ],
    'Feed': [
        ('FEED-001', 'Load feed with available tasks', 'User logged in, Navigate to Feed page', 'All available tasks displayed in list/card format, Task title, description, and status visible, Pagination controls shown if tasks > 10', '', ''),
        ('FEED-002', 'Load feed with no available tasks', 'User logged in, No tasks in database, Navigate to Feed page', 'Empty state message displayed: No tasks available, Call-to-action button or message shown', '', ''),
        ('FEED-003', 'View task details', 'User on Feed page, Click on specific task card', 'Modal/detail page opens, Task title, description, attached documents, budget, deadline all displayed, Apply button available', '', ''),
        ('FEED-004', 'Filter tasks by category', 'User on Feed page, Select category filter: Development', 'Feed refreshes, Only tasks with Development category displayed, Filter tag shown as active', '', ''),
        ('FEED-005', 'Filter tasks by status', 'User on Feed page, Select status filter: Open', 'Feed refreshes, Only open tasks displayed, Task count updated, Other statuses hidden', '', ''),
        ('FEED-006', 'Pagination - Next page', 'User on Feed page, 15+ tasks available, Click Next button', 'Second set of 10 tasks displayed, Page indicator updated, Previous button enabled', '', ''),
        ('FEED-007', 'Pagination - Previous page', 'User on page 2 of Feed, Click Previous button', 'First page of tasks displayed, Page indicator updated to page 1, Next button enabled', '', ''),
        ('FEED-008', 'Search tasks by keyword', 'User on Feed page, Search: logo design', 'Feed filtered to show matching tasks, Result count displayed, No results message if none match', '', ''),
        ('FEED-009', 'Scroll to load more tasks', 'User on Feed page, Scroll to bottom of list', 'Additional tasks loaded automatically, Loading indicator shown briefly, New tasks appended to list', '', ''),
        ('FEED-010', 'Apply for task', 'User viewing task details, Click Apply button', 'Application submitted successfully, Confirmation message displayed, Task moved to My Requests, Status changed to Applied', '', ''),
    ],
    'My Task': [
        ('TASK-001', 'View assigned tasks list', 'User logged in, User has assigned tasks, Navigate to My Tasks', 'All assigned tasks displayed in list/card format, Task title, status, deadline, assignee visible', '', ''),
        ('TASK-002', 'View assigned tasks with no tasks', 'User logged in, User has no assigned tasks, Navigate to My Tasks', 'Empty state message displayed: No tasks assigned yet, Call-to-action or message shown', '', ''),
        ('TASK-003', 'View task status - In Progress', 'User on My Tasks page, View task with status In Progress', 'Status badge displayed with In Progress label, Progress percentage shown (if available), Color-coded indicator visible', '', ''),
        ('TASK-004', 'View task status - Completed', 'User on My Tasks page, View task with status Completed', 'Status badge displayed with Completed label, Completion checkmark shown, Green color indicator', '', ''),
        ('TASK-005', 'View complete task details', 'User on My Tasks page, Click on specific task', 'Task detail page/modal opens, Task description, attachments, budget, deadline, assigned by, notes all displayed', '', ''),
        ('TASK-006', 'Update task status', 'User viewing task details, Change status to In Progress, Click Save', 'Status updated successfully, Confirmation message shown, Task list refreshed with new status', '', ''),
        ('TASK-007', 'Add task comment', 'User viewing task details, Enter comment: Starting work on this, Click Add Comment', 'Comment added successfully, Timestamp displayed, Comment visible in task feed', '', ''),
        ('TASK-008', 'Attach file to task', 'User viewing task details, Click Attach File, Select file (max 5MB), Click Upload', 'File uploaded successfully, File name displayed in attachments section, File preview available', '', ''),
        ('TASK-009', 'Refresh tasks list', 'User on My Tasks page, Click Refresh button or pull down', 'Tasks list refreshed, Latest status and data loaded, Timestamp updated', '', ''),
        ('TASK-010', 'Mark task as complete', 'User viewing task details, Click Mark Complete button', 'Status changed to Completed, Completion date recorded, Task moved to completed section, Notification sent to task creator', '', ''),
    ],
    'My Request': [
        ('REQ-001', 'Create new request', 'User on Feed, Found suitable task, Click Apply, Provide proposal: I can complete this in 5 days', 'Request created successfully, Confirmation message displayed, Request visible in My Requests with Pending status', '', ''),
        ('REQ-002', 'View all created requests', 'User logged in, Has 3+ requests, Navigate to My Requests', 'All requests displayed in list format, Request ID, task name, status, date applied all visible', '', ''),
        ('REQ-003', 'View request status - Pending', 'User on My Requests page, View request with Pending status', 'Status badge shows Pending with yellow indicator, Created date displayed, Cancel Request option available', '', ''),
        ('REQ-004', 'View request status - Accepted', 'User on My Requests page, View request with Accepted status', 'Status badge shows Accepted with green indicator, Acceptance date displayed, View Task button available', '', ''),
        ('REQ-005', 'View request status - Rejected', 'User on My Requests page, View request with Rejected status', 'Status badge shows Rejected with red indicator, Rejection date displayed, Rejection reason shown if provided, Apply Again option available', '', ''),
        ('REQ-006', 'Cancel pending request', 'User on My Requests page, View pending request, Click Cancel Request', 'Confirmation dialog displayed, Upon confirmation, Request status changed to Cancelled, Cancellation date recorded', '', ''),
        ('REQ-007', 'View request details', 'User on My Requests page, Click on specific request card', 'Request detail page opens, Task details, proposal text, timeline, status, requester info, and conversation thread visible', '', ''),
        ('REQ-008', 'Multiple requests on single task', 'User on Feed, Same task has 5+ applications, View task details', 'Request counter shows 5 applications, All requests listed with status and applicant info, Latest requests shown first', '', ''),
        ('REQ-009', 'View request conversation', 'User on My Requests page, Open request with messages, View conversation', 'All messages between user and task creator displayed in chronological order, Timestamps visible, Latest messages at bottom', '', ''),
        ('REQ-010', 'Filter requests by status', 'User on My Requests page, Filter by Accepted', 'List refreshed showing only accepted requests, Filter tag displayed as active, Request count updated', '', ''),
    ],
    'Notification': [
        ('NOT-001', 'Receive task assignment notification', 'Admin assigns task to user, User is logged in', 'Notification badge appears (number indicator), Toast notification shown at top/bottom, Sound alert plays (if enabled)', '', ''),
        ('NOT-002', 'Receive request status change notification', 'Request status changes from Pending to Accepted', 'Notification badge updated, Notification displayed with status change details, Green checkmark indicator', '', ''),
        ('NOT-003', 'View notification list', 'User logged in, Click Notification bell icon, 5+ unread notifications exist', 'Notification list sidebar/panel opens, All notifications displayed with most recent first, Unread notifications highlighted/bold', '', ''),
        ('NOT-004', 'Mark notification as read', 'User viewing notification list, Click on unread notification', 'Notification marked as read, Highlight/bold removed, Read indicator shown, Badge count decremented', '', ''),
        ('NOT-005', 'Mark all notifications as read', 'User viewing notification list with 5+ unread, Click Mark All as Read', 'All notifications marked as read, Highlight removed from all items, Badge count reset to 0', '', ''),
        ('NOT-006', 'No notifications state', 'User logged in, No notifications available, Click Notification bell icon', 'Empty state message displayed: No notifications, Icon shows 0 count, Sidebar displays friendly message', '', ''),
        ('NOT-007', 'Click notification to navigate', 'User viewing notification for task assignment, Click notification', 'User navigated to task details page, Task information fully loaded, Notification marked as read', '', ''),
        ('NOT-008', 'Click notification to navigate to request', 'User viewing notification for request status update, Click notification', 'User navigated to My Requests page, Specific request highlighted/focused, Notification marked as read', '', ''),
        ('NOT-009', 'Delete/Dismiss notification', 'User on notification list, Hover over notification, Click Delete/X button', 'Notification removed from list, Confirmation message shown (optional), List refreshed', '', ''),
        ('NOT-010', 'Real-time notification update', 'Two users: one creates request for another task, Second user is viewing dashboard', 'New notification appears in real-time (within 2-3 seconds), No page refresh required, Badge count updated immediately, Toast notification shown', '', ''),
    ]
}

# Define styles
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
alignment_wrap = Alignment(wrap_text=True, vertical='top')

# Create sheets for each module
for module_name, cases in test_cases.items():
    ws = wb.create_sheet(module_name)
    
    # Add headers
    headers = ['Test case No', 'Test case Scenario', 'Test Input', 'Expected Output', 'Actual Output', 'Result (pass/fail)']
    ws.append(headers)
    
    # Format header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Add test cases
    for case in cases:
        ws.append(case)
    
    # Format data rows
    for row in ws.iter_rows(min_row=2, max_row=len(cases) + 1):
        for cell in row:
            cell.border = border
            cell.alignment = alignment_wrap
    
    # Set column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 15
    
    # Set header row height
    ws.row_dimensions[1].height = 30
    
    # Add data validation for Result column (dropdown with Pass/Fail)
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Pass,Fail"', allow_blank=False)
    dv.error = 'You must select Pass or Fail'
    dv.errorTitle = 'Invalid Entry'
    dv.prompt = 'Select Pass or Fail'
    dv.promptTitle = 'Result'
    ws.add_data_validation(dv)
    
    # Apply dropdown to Result column (column F, rows 2 onwards)
    for row in range(2, len(cases) + 2):
        dv.add(ws[f'F{row}'])

# Save workbook
output_path = 'HireHelper_TestCases.xlsx'
wb.save(output_path)
print(f'Excel file created successfully: {output_path}')
